"""Bounded XLSX and UTF-8 delimited readers; no formula evaluation."""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterator
from contextlib import suppress
from datetime import date, datetime, time
from itertools import islice, zip_longest
from pathlib import PurePath
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import fromstring
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from obox_builder.conf import settings


class ImportFileError(ValueError):
    pass


def check_file(filename: str, content: bytes) -> str:
    suffix = PurePath(filename).suffix.lower()
    if suffix not in {".xlsx", ".csv", ".tsv"}:
        raise ImportFileError("Choose an XLSX, CSV or TSV file.")
    if not content or len(content) > settings.IMPORT_MAX_FILE_BYTES:
        raise ImportFileError("The file is empty or exceeds the upload limit.")
    return suffix


def _check_archive(content: bytes) -> None:
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            names = [entry.filename for entry in entries]
            if len(entries) > 2000 or len(names) != len(set(names)):
                raise ImportFileError("Invalid or oversized workbook archive.")
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ImportFileError("The file is not an XLSX workbook.")
            total = 0
            for entry in entries:
                name = entry.filename.lower()
                if entry.flag_bits & 1 or "vbaproject" in name or "externallinks/" in name:
                    raise ImportFileError(
                        "Encrypted workbooks, macros and external links are unsupported."
                    )
                total += entry.file_size
                if total > settings.IMPORT_MAX_UNPACKED_BYTES:
                    raise ImportFileError("The expanded workbook exceeds the size limit.")
            # Read through a bounded stream rather than trusting ZIP metadata alone.
            remaining = settings.IMPORT_MAX_UNPACKED_BYTES
            for entry in entries:
                with archive.open(entry) as stream:
                    while block := stream.read(min(65536, remaining + 1)):
                        remaining -= len(block)
                        if remaining < 0:
                            raise ImportFileError("The expanded workbook exceeds the size limit.")
            types = fromstring(archive.read("[Content_Types].xml"))
            if any("macroEnabled" in item.attrib.get("ContentType", "") for item in types):
                raise ImportFileError("Macro-enabled workbooks are unsupported.")
    except (BadZipFile, KeyError, RuntimeError) as exc:
        raise ImportFileError("The workbook archive is invalid.") from exc


def _cell(cell: Any, cached: Any) -> Any:
    if cell.data_type == "f":
        if cached.value is None:
            return {"error": "Formula has no cached value. Save calculated values in Excel first."}
        cell = cached
    if cell.data_type == "e":
        return {"error": "The Excel cell contains an error."}
    value = cell.value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and value == int(value)
        and re.fullmatch(r"0{2,}", cell.number_format or "")
    ):
        return str(int(value)).zfill(len(cell.number_format))
    return value


class Spreadsheet:
    def __init__(self, filename: str, content: bytes) -> None:
        self.suffix = check_file(filename, content)
        self.content = content
        self.workbook: Any = None
        self.cached: Any = None

    def __enter__(self) -> Spreadsheet:
        if self.suffix == ".xlsx":
            try:
                _check_archive(self.content)
                self.workbook = load_workbook(
                    io.BytesIO(self.content), read_only=True, keep_links=False
                )
                self.cached = load_workbook(
                    io.BytesIO(self.content), read_only=True, data_only=True, keep_links=False
                )
            except ImportFileError:
                raise
            except (
                InvalidFileException,
                BadZipFile,
                DefusedXmlException,
                ParseError,
                KeyError,
                ValueError,
            ) as exc:
                self.__exit__(None, None, None)
                raise ImportFileError("The workbook could not be read safely.") from exc
        return self

    def __exit__(self, *args: Any) -> None:
        for workbook in (self.workbook, self.cached):
            if workbook is not None:
                workbook.close()

    @property
    def sheets(self) -> list[str]:
        if self.workbook is None:
            return ["Data"]
        return [sheet.title for sheet in self.workbook if sheet.sheet_state == "visible"]

    def rows(self, sheet: str) -> Iterator[tuple[int, list[Any]]]:
        try:
            yield from self._rows(sheet)
        except (csv.Error, BadZipFile, DefusedXmlException, ParseError) as exc:
            raise ImportFileError("The worksheet could not be read safely.") from exc

    def _rows(self, sheet: str) -> Iterator[tuple[int, list[Any]]]:
        if sheet not in self.sheets:
            raise ImportFileError("Choose a visible worksheet.")
        if self.workbook is None:
            try:
                text = self.content.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise ImportFileError("Save the delimited file as UTF-8.") from exc
            if "\x00" in text:
                raise ImportFileError("The delimited file contains binary data.")
            dialect: Any = csv.excel_tab if self.suffix == ".tsv" else csv.excel
            with suppress(csv.Error):
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            if dialect is csv.excel:
                with suppress(csv.Error):
                    dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t")
            rows = csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True)
            for number, row in enumerate(rows, 1):
                yield number, self._bounded(row)
            return
        original, cached = self.workbook[sheet], self.cached[sheet]
        original.reset_dimensions()
        cached.reset_dimensions()
        for number, (cells, values) in enumerate(
            zip_longest(original.iter_rows(), cached.iter_rows(), fillvalue=()), 1
        ):
            if len(cells) > settings.IMPORT_MAX_COLUMNS:
                raise ImportFileError("The worksheet exceeds the column limit.")
            yield (
                number,
                self._bounded(
                    [_cell(cell, value) for cell, value in zip(cells, values, strict=True)]
                ),
            )

    @staticmethod
    def _bounded(row: list[Any]) -> list[Any]:
        if len(row) > settings.IMPORT_MAX_COLUMNS:
            raise ImportFileError("The worksheet exceeds the column limit.")
        if any(isinstance(value, str) and len(value) > 65536 for value in row):
            raise ImportFileError("A cell exceeds the text length limit.")
        return row

    def sample(self, sheet: str) -> list[list[Any]]:
        return [row for _, row in islice(self.rows(sheet), 10)]
