from __future__ import annotations

import io
from datetime import date
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from obox_builder.conf import settings
from obox_builder.contrib.imports.parser import ImportFileError, Spreadsheet
from openpyxl import Workbook


def workbook_bytes(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


@pytest.mark.usefixtures("imports_env")
@pytest.mark.parametrize(
    "filename,content,expected",
    [
        (
            "test.csv",
            '\ufeffname;note\r\nTiếng Việt;"first\nsecond"\r\n'.encode(),
            [["name", "note"], ["Tiếng Việt", "first\nsecond"]],
        ),
        ("test.tsv", b"name\tnote\nA\tB\n", [["name", "note"], ["A", "B"]]),
        ("test.csv", b"reference\nA\n", [["reference"], ["A"]]),
    ],
)
def test_delimited_files(filename: str, content: bytes, expected: list[list[str]]) -> None:
    with Spreadsheet(filename, content) as spreadsheet:
        assert spreadsheet.sheets == ["Data"]
        assert [row for _, row in spreadsheet.rows("Data")] == expected


@pytest.mark.usefixtures("imports_env")
def test_excel_dates_hidden_sheets_formulas_and_zero_padding() -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["code", "date", "formula"])
    sheet.append([12, date(2026, 10, 23), "=1+1"])
    sheet["A2"].number_format = "00000"
    workbook.create_sheet("Hidden").sheet_state = "hidden"
    output = io.BytesIO()
    workbook.save(output)
    with Spreadsheet("data.xlsx", output.getvalue()) as spreadsheet:
        assert spreadsheet.sheets == [sheet.title]
        rows = list(spreadsheet.rows(sheet.title))
        assert rows[1][1][:2] == ["00012", "2026-10-23T00:00:00"]
        assert "cached" in rows[1][1][2]["error"]
        with pytest.raises(ImportFileError, match="visible"):
            list(spreadsheet.rows("Hidden"))


@pytest.mark.usefixtures("imports_env")
@pytest.mark.parametrize(
    "filename,content",
    [
        ("file.xlsm", b"invalid"),
        ("file.xls", b"invalid"),
        ("file.xlsx", b"PKbogus"),
        ("file.csv", b"\xff\xfe"),
        ("file.csv", b"a\x00b"),
        ("file.csv", b""),
    ],
)
def test_invalid_files_are_rejected(filename: str, content: bytes) -> None:
    with pytest.raises(ImportFileError), Spreadsheet(filename, content) as spreadsheet:
        list(spreadsheet.rows(spreadsheet.sheets[0]))


@pytest.mark.usefixtures("imports_env")
def test_archive_limits_macros_external_links_and_xml_entities(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = workbook_bytes([["name"], ["value"]])
    for entry in ("xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"):
        output = io.BytesIO(source)
        with ZipFile(output, "a", ZIP_DEFLATED) as archive:
            archive.writestr(entry, b"unsafe")
        with pytest.raises(ImportFileError), Spreadsheet("file.xlsx", output.getvalue()):
            pass
    monkeypatch.setattr(settings.load(), "IMPORT_MAX_UNPACKED_BYTES", 10)
    with pytest.raises(ImportFileError, match="expanded"), Spreadsheet("file.xlsx", source):
        pass


@pytest.mark.usefixtures("imports_env")
def test_column_and_cell_limits(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings.load(), "IMPORT_MAX_COLUMNS", 1)
    with (
        pytest.raises(ImportFileError, match="column"),
        Spreadsheet("file.csv", b"a,b\n") as spreadsheet,
    ):
        list(spreadsheet.rows("Data"))
    with (
        pytest.raises(ImportFileError, match="length"),
        Spreadsheet("file.csv", ("a" * 65537).encode()) as spreadsheet,
    ):
        list(spreadsheet.rows("Data"))
